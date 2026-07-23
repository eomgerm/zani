"""Prepare an official-style EngageNet contract using NTFS hard links only."""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import os
import re
import sys
import tempfile
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from zani_ai.engagement.contracts import LABELS, load_dataset_contract


SPLITS = {
    "Train": ("train", "train_engagement_labels.xlsx"),
    "Validation": ("valid", "validation_engagement_labels.xlsx"),
    "Test": ("test", "test_engagement_labels.csv"),
}
SNP_LABEL = "SNP(Subject Not Present)"
SUBJECT_PATTERN = re.compile(r"^subject_(?P<subject>[^_]+)_.+\.mp4$", re.IGNORECASE)
LABEL_LOOKUP = {label.casefold(): label for label in LABELS}
LABEL_LOOKUP["barely-engaged"] = "Barely-Engaged"


@dataclass(frozen=True, slots=True)
class Clip:
    split_name: str
    contract_split: str
    chunk: str
    clip_id: str
    label: str
    subject_id: str
    source_video: Path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-root", required=True, type=Path)
    parser.add_argument("--output-root", required=True, type=Path)
    return parser.parse_args()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for block in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _read_xlsx(path: Path) -> tuple[list[dict[str, object]], set[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:  # pragma: no cover - environment-dependent
        raise RuntimeError("XLSX labels require openpyxl; install the 'eda' extra first.") from error
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if headers is None:
            return [], set()
        names = ["" if value is None else str(value).strip() for value in headers]
        return [dict(zip(names, row, strict=False)) for row in rows], set(names)
    finally:
        workbook.close()


def read_label_table(path: Path) -> list[dict[str, object]]:
    if not path.is_file():
        raise FileNotFoundError(f"missing label file: {path}")
    if path.suffix.casefold() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as file:
            reader = csv.DictReader(file)
            rows = list(reader)
            columns = set(reader.fieldnames or ())
    else:
        rows, columns = _read_xlsx(path)
    if not {"chunk", "label"}.issubset(columns):
        raise ValueError(f"{path} requires columns 'chunk' and 'label'")
    return rows


def inside(root: Path, path: Path) -> bool:
    try:
        path.resolve().relative_to(root)
    except ValueError:
        return False
    return True


def source_fingerprint(path: Path) -> dict[str, int]:
    stat = path.stat()
    return {"device": stat.st_dev, "inode": stat.st_ino, "size": stat.st_size}


def collect_clips(source_root: Path) -> tuple[list[Clip], dict[str, Any]]:
    clips: list[Clip] = []
    seen_clip_ids: set[str] = set()
    subject_splits: dict[str, set[str]] = defaultdict(set)
    raw_counts: dict[str, Counter[str]] = {}
    excluded: list[dict[str, str]] = []
    label_hashes: dict[str, str] = {}
    label_metadata: dict[str, dict[str, object]] = {}

    for split_name, (contract_split, label_file) in SPLITS.items():
        split_dir = source_root / split_name
        if not split_dir.is_dir():
            raise FileNotFoundError(f"missing source video directory: {split_dir}")
        labels_path = source_root / label_file
        rows = read_label_table(labels_path)
        label_hashes[label_file] = sha256_file(labels_path)
        stat = labels_path.stat()
        label_metadata[label_file] = {
            "path": str(labels_path.resolve()),
            "bytes": stat.st_size,
            "modified_at": datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(),
        }
        counts: Counter[str] = Counter()
        for number, row in enumerate(rows, start=2):
            chunk = str(row.get("chunk", "")).strip()
            raw_label = str(row.get("label", "")).strip()
            match = SUBJECT_PATTERN.fullmatch(chunk)
            if match is None:
                raise ValueError(f"{labels_path.name} row {number} has malformed chunk: {chunk!r}")
            if Path(chunk).name != chunk:
                raise ValueError(f"{labels_path.name} row {number} has unsafe chunk path: {chunk!r}")
            clip_id = Path(chunk).stem
            if not clip_id or clip_id in seen_clip_ids:
                raise ValueError(f"clip IDs must be globally unique; duplicate or empty: {clip_id!r}")
            seen_clip_ids.add(clip_id)
            label = LABEL_LOOKUP.get(raw_label.casefold())
            if raw_label != SNP_LABEL and label is None:
                raise ValueError(f"{labels_path.name} row {number} has unknown label: {raw_label!r}")
            counts[raw_label if raw_label == SNP_LABEL else label] += 1
            source_video = (split_dir / chunk).resolve()
            if not inside(split_dir.resolve(), source_video):
                raise ValueError(f"source video escapes its split directory: {chunk!r}")
            if not source_video.is_file():
                raise FileNotFoundError(f"missing source video: {source_video}")
            subject_id = match.group("subject")
            subject_splits[subject_id].add(contract_split)
            if raw_label == SNP_LABEL:
                excluded.append({"clip_id": clip_id, "split": contract_split, "subject_id": subject_id})
                continue
            clips.append(
                Clip(split_name, contract_split, chunk, clip_id, label, subject_id, source_video)
            )
        raw_counts[contract_split] = Counter({label: counts[label] for label in (*LABELS, SNP_LABEL)})

    overlap = {subject: sorted(names) for subject, names in subject_splits.items() if len(names) > 1}
    if overlap:
        raise ValueError(f"subjects appear in multiple splits: {overlap}")
    return clips, {
        "label_file_sha256": label_hashes,
        "label_file_metadata": label_metadata,
        "raw_counts_by_split_and_label": {key: dict(value) for key, value in raw_counts.items()},
        "excluded_snp_clips": excluded,
        "subject_counts_by_split": {
            split: len({clip.subject_id for clip in clips if clip.contract_split == split})
            for split in ("train", "valid", "test")
        },
    }


def atomic_write(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile("w", encoding="utf-8", newline="", dir=path.parent, delete=False) as file:
        temporary = Path(file.name)
        file.write(content)
    try:
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def write_metadata(output_root: Path, clips: list[Clip], manifest: dict[str, Any]) -> None:
    labels_file = io.StringIO(newline="")
    writer = csv.writer(labels_file, lineterminator="\n")
    writer.writerow(("clip_id", "label", "subject_id"))
    writer.writerows((clip.clip_id, clip.label, clip.subject_id) for clip in clips)
    atomic_write(output_root / "final_labels.csv", labels_file.getvalue())
    for split in ("train", "valid", "test"):
        clip_ids = [clip.clip_id for clip in clips if clip.contract_split == split]
        atomic_write(output_root / f"{split}.txt" if split != "valid" else output_root / "valid.txt", "\n".join(clip_ids) + "\n")
    atomic_write(output_root / "preparation_manifest.json", json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")


def verify_existing_hard_link(source: Path, destination: Path) -> None:
    if destination.exists():
        source_identity = source_fingerprint(source)
        destination_identity = source_fingerprint(destination)
        if (
            destination_identity["size"] != source_identity["size"]
            or destination_identity != source_identity
            or not os.path.samefile(source, destination)
        ):
            raise FileExistsError(
                f"existing destination is not the same hard-linked source (size/fingerprint mismatch): {destination}"
            )


def ensure_hard_link(source: Path, destination: Path) -> None:
    if destination.exists():
        verify_existing_hard_link(source, destination)
        return
    destination.parent.mkdir(parents=True, exist_ok=True)
    try:
        os.link(source, destination)
    except OSError as error:
        raise RuntimeError(
            f"could not create required NTFS hard link {destination} -> {source}; copying and symlinks are disabled"
        ) from error


def main() -> int:
    args = parse_args()
    source_root = args.source_root.resolve()
    output_root = args.output_root.resolve()
    if not source_root.is_dir():
        raise FileNotFoundError(f"source root does not exist: {source_root}")
    if output_root == source_root or inside(source_root, output_root):
        raise ValueError("output root must not be inside source root; source data must remain untouched")
    clips, collected = collect_clips(source_root)
    clips.sort(key=lambda clip: clip.clip_id)
    videos_root = output_root / "videos"
    if not inside(output_root, videos_root):
        raise ValueError("videos output path escapes output root")
    included_counts: dict[str, dict[str, int]] = {}
    for split in ("train", "valid", "test"):
        counts = Counter(clip.label for clip in clips if clip.contract_split == split)
        included_counts[split] = {label: counts[label] for label in LABELS}
    manifest: dict[str, Any] = {
        "format": "engagenet_contract_v1",
        "prepared_at": datetime.now(timezone.utc).isoformat(),
        "source": {"root": str(source_root), "splits": {key: str(source_root / key) for key in SPLITS}},
        "source_file_metadata": collected["label_file_metadata"],
        "source_file_sha256": collected["label_file_sha256"],
        "raw_counts_by_split_and_label": collected["raw_counts_by_split_and_label"],
        "included_counts_by_split_and_label": included_counts,
        "excluded_snp_clips": collected["excluded_snp_clips"],
        "subject_counts_by_split": collected["subject_counts_by_split"],
        "link_strategy": {"type": "ntfs_hard_link", "copy_fallback": False, "symlink_fallback": False},
    }
    link_targets: list[tuple[Clip, Path]] = []
    for clip in clips:
        destination = videos_root / f"{clip.clip_id}.mp4"
        if not inside(output_root, destination):
            raise ValueError(f"video output path escapes output root: {destination}")
        link_targets.append((clip, destination))

    # Do not replace a previously valid metadata contract until every existing
    # destination has been proven to be the expected hard link.
    for clip, destination in link_targets:
        verify_existing_hard_link(clip.source_video, destination)

    total = len(clips)
    for index, (clip, destination) in enumerate(link_targets, start=1):
        ensure_hard_link(clip.source_video, destination)
        if index == total or index % 500 == 0:
            print(f"Linked {index:,}/{total:,} videos", flush=True)
    write_metadata(output_root, clips, manifest)
    load_dataset_contract(output_root)
    print(f"Prepared valid contract: {output_root}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
